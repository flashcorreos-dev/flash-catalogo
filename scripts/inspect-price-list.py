from pathlib import Path
import json
import openpyxl

source = Path('/home/ubuntu/upload/LISTAJUNIO2026m..xlsx')
book = openpyxl.load_workbook(source, data_only=False)
result = {"sheets": book.sheetnames, "details": []}
for sheet in book.worksheets:
    detail = {"title": sheet.title, "max_row": sheet.max_row, "max_column": sheet.max_column, "rows": []}
    for row_number in list(range(1, 8)) + [264, 265, 266, 267, 268]:
        values = [sheet.cell(row_number, col).value for col in range(1, min(sheet.max_column, 12) + 1)]
        detail["rows"].append({"row": row_number, "values": values})
    detail["merged_ranges"] = [str(item) for item in list(sheet.merged_cells.ranges)[:20]]
    result["details"].append(detail)
print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
